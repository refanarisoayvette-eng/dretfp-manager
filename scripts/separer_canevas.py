"""
Script pour séparer le fichier FUSION_CANEVAS_TPA_2027.xlsx
en plusieurs fichiers, un par établissement.

Utilisation :
    python scripts/separer_canevas.py
"""

import os
import openpyxl
from openpyxl import Workbook
from copy import copy


# ============================================================
# CONFIGURATION
# ============================================================

FICHIER_SOURCE = 'donnees/source/FUSION_CANEVAS_TPA_2027.xlsx'
DOSSIER_SORTIE = 'donnees/canevas'

# Correspondance : nom d'établissement → liste de noms de feuilles à copier
ETABLISSEMENTS = {
    'CFP_AMBOSITRA': {
        'nom_affiche': 'CFP AMBOSITRA',
        'feuilles_source': ['CFP AMBOSITRA'],
    },
    'CFPF_FANDRIANA': {
        'nom_affiche': 'CFPF FANDRIANA',
        'feuilles_source': ['CFPF FANDRIANA'],
    },
    'CFP_FIADANANA_FANDRIANA': {
        'nom_affiche': 'CFP FIADANANA FANDRIANA',
        'feuilles_source': ['DRETFP ET SES ETABLISSEMENTS (6'],
    },
    'LTP_AMBOSITRA': {
        'nom_affiche': 'LTP AMBOSITRA',
        'feuilles_source': ['DRETFP ET SES ETABLISSEMENTS (2'],
    },
    'LTP_MIARINAVARATRA': {
        'nom_affiche': 'LTP MIARINAVARATRA',
        'feuilles_source': [
            'DRETFP ET SES ETABLISSEMENTS (3',
            'DRETFP ET SES ETABLISSEMENTS (5',
        ],
    },
    'LTPA_FANDRIANA': {
        'nom_affiche': 'LTPA Fandriana',
        'feuilles_source': ['DRETFP ET SES ETABLISSEMENTS (1'],
    },
    'LTPA_AMBINDA_FANDRIANA': {
        'nom_affiche': 'LTPA Ambinda Fandriana',
        'feuilles_source': ['ORGANISMES RATTACHES (1'],
    },
}


# ============================================================
# FONCTIONS
# ============================================================

def copier_feuille(ws_source, wb_dest):
    """Copie une feuille complète dans un nouveau classeur."""
    ws_dest = wb_dest.create_sheet(title=ws_source.title[:31])
    
    # Copier les cellules
    for row in ws_source.iter_rows():
        for cell in row:
            nouvelle_cellule = ws_dest.cell(
                row=cell.row,
                column=cell.column,
                value=cell.value
            )
            # Copier le style
            if cell.has_style:
                try:
                    nouvelle_cellule.font = copy(cell.font)
                    nouvelle_cellule.border = copy(cell.border)
                    nouvelle_cellule.fill = copy(cell.fill)
                    nouvelle_cellule.number_format = cell.number_format
                    nouvelle_cellule.alignment = copy(cell.alignment)
                except:
                    pass
    
    # Copier la largeur des colonnes
    for col_letter, dim in ws_source.column_dimensions.items():
        ws_dest.column_dimensions[col_letter].width = dim.width
    
    return ws_dest


def creer_fichier_etablissement(nom_etab, config, wb_source):
    """Crée un fichier Excel pour un établissement."""
    nom_fichier = f"{nom_etab}.xlsx"
    chemin_sortie = os.path.join(DOSSIER_SORTIE, nom_fichier)
    
    print(f"\n📄 Création : {nom_fichier}")
    print(f"   Établissement : {config['nom_affiche']}")
    
    wb_dest = Workbook()
    # Supprimer la feuille par défaut
    if 'Sheet' in wb_dest.sheetnames:
        wb_dest.remove(wb_dest['Sheet'])
    
    nb_feuilles = 0
    for nom_feuille in config['feuilles_source']:
        if nom_feuille not in wb_source.sheetnames:
            print(f"   ⚠️  Feuille '{nom_feuille}' introuvable, ignorée.")
            continue
        
        ws_source = wb_source[nom_feuille]
        copier_feuille(ws_source, wb_dest)
        nb_feuilles += 1
        print(f"   ✓ Feuille copiée : {nom_feuille} ({ws_source.max_row} lignes)")
    
    if nb_feuilles == 0:
        print(f"   ❌ Aucune feuille copiée, fichier ignoré.")
        return False
    
    wb_dest.save(chemin_sortie)
    print(f"   💾 Sauvegardé : {chemin_sortie}")
    return True


# ============================================================
# PROGRAMME PRINCIPAL
# ============================================================

def main():
    print("=" * 60)
    print("  SÉPARATION DU FICHIER FUSION EN 7 FICHIERS")
    print("=" * 60)
    
    # Vérifier que le fichier source existe
    if not os.path.exists(FICHIER_SOURCE):
        print(f"\n❌ ERREUR : '{FICHIER_SOURCE}' n'existe pas.")
        print(f"   Placez le fichier dans : donnees/source/")
        return
    
    # Créer le dossier de sortie si nécessaire
    os.makedirs(DOSSIER_SORTIE, exist_ok=True)
    
    # Ouvrir le fichier source
    print(f"\n📂 Ouverture de : {FICHIER_SOURCE}")
    wb_source = openpyxl.load_workbook(FICHIER_SOURCE, data_only=False)
    print(f"   {len(wb_source.sheetnames)} feuilles trouvées :")
    for nom in wb_source.sheetnames:
        print(f"     - {nom}")
    
    # Créer un fichier par établissement
    print(f"\n{'─' * 60}")
    print("  CRÉATION DES FICHIERS PAR ÉTABLISSEMENT")
    print(f"{'─' * 60}")
    
    nb_crees = 0
    for nom_etab, config in ETABLISSEMENTS.items():
        if creer_fichier_etablissement(nom_etab, config, wb_source):
            nb_crees += 1
    
    # Résumé
    print(f"\n{'=' * 60}")
    print(f"  RÉSUMÉ")
    print(f"{'=' * 60}")
    print(f"  ✅ Fichiers créés : {nb_crees} / {len(ETABLISSEMENTS)}")
    print(f"  📁 Dossier : {DOSSIER_SORTIE}")
    print(f"\n✅ SÉPARATION TERMINÉE !")
    
    # Lister les fichiers créés
    print(f"\n📄 Fichiers dans {DOSSIER_SORTIE} :")
    for f in sorted(os.listdir(DOSSIER_SORTIE)):
        if f.endswith('.xlsx'):
            taille = os.path.getsize(os.path.join(DOSSIER_SORTIE, f))
            print(f"   - {f} ({taille // 1024} Ko)")


if __name__ == '__main__':
    main()