"""
Diagnostic pour trouver les colonnes de montants dans les fichiers Excel.
"""

import os
import glob
import openpyxl


DOSSIER_CANEVAS = 'donnees/canevas'


def diagnostiquer(chemin):
    nom = os.path.basename(chemin)
    print("")
    print("=" * 70)
    print("  FICHIER : " + nom)
    print("=" * 70)
    
    wb = openpyxl.load_workbook(chemin, data_only=True)
    
    for nom_feuille in wb.sheetnames:
        ws = wb[nom_feuille]
        print("")
        print("  Feuille : '" + nom_feuille + "'")
        
        # Lire les 3 premières lignes
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True), start=1):
            if row:
                valeurs = []
                for j, v in enumerate(row[:15]):
                    if v is not None and str(v).strip():
                        valeurs.append("Col" + str(j) + "='" + str(v)[:20] + "'")
                print("    Ligne " + str(i) + " : " + " | ".join(valeurs))


def main():
    print("=" * 70)
    print("  DIAGNOSTIC DES MONTANTS")
    print("=" * 70)
    
    fichiers = sorted(glob.glob(os.path.join(DOSSIER_CANEVAS, '*.xlsx')))
    for f in fichiers:
        diagnostiquer(f)


if __name__ == '__main__':
    main()