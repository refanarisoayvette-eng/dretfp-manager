import os
import sys
import django
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

FICHIER_EXCEL = 'FUSION_CANEVAS_TPA_2027.xlsx'

wb = openpyxl.load_workbook(FICHIER_EXCEL, data_only=True)

print("=" * 70)
print("  DIAGNOSTIC DES FEUILLES")
print("=" * 70)

for nom in wb.sheetnames:
    ws = wb[nom]
    print(f"\n📄 Feuille : '{nom}' ({ws.max_row} lignes)")
    print("-" * 70)
    
    # Afficher les 3 premières lignes
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True), start=1):
        valeurs = [str(v)[:20] if v else '' for v in row[:13]]
        print(f"  Ligne {i}: {valeurs}")