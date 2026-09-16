"""
Importation des canevas par etablissement.
Version finale - gere les montants + produits dans sous_activite.
Sans accents pour eviter les problemes d'encodage.
"""

import os
import sys
import glob
import django
import openpyxl
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

from core.models import (
    Etablissement, Programme, Activite, SousActivite,
    Produit, PCOP, LigneBudgetaire
)


DOSSIER_CANEVAS = 'donnees/canevas'


def nom_fichier_vers_etab(nom_fichier):
    base = os.path.splitext(os.path.basename(nom_fichier))[0]
    return base.replace('_', ' ')


def detecter_fin_donnees(ws):
    """Detecte la derniere ligne avec des donnees reelles."""
    derniere = 1
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row and any(v is not None and str(v).strip() for v in row):
            derniere = row_num
    return derniere


def importer_fichier(chemin):
    nom_fichier = os.path.basename(chemin)
    nom_etab = nom_fichier_vers_etab(chemin)
    
    print("")
    print("=" * 60)
    print("  FICHIER : " + nom_fichier)
    print("  ETABLISSEMENT : " + nom_etab)
    print("=" * 60)
    
    type_etab = 'LTP'
    upper = nom_etab.upper()
    if 'CFPF' in upper:
        type_etab = 'CFPF'
    elif 'CFP' in upper:
        type_etab = 'CFP'
    elif 'LTPA' in upper:
        type_etab = 'LTPA'
    
    etab, created = Etablissement.objects.get_or_create(
        nom=nom_etab,
        defaults={
            'type_etablissement': type_etab,
            'district': "Amoron'i Mania",
        }
    )
    if created:
        print("  + Etablissement cree : " + nom_etab)
    
    wb = openpyxl.load_workbook(chemin, data_only=True)
    
    total_produits = 0
    total_lignes = 0
    total_ignorees = 0
    total_montants = 0
    
    for nom_feuille in wb.sheetnames:
        ws = wb[nom_feuille]
        derniere_ligne = detecter_fin_donnees(ws)
        
        print("")
        print("  Feuille : '" + nom_feuille + "'")
        print("    Fin des donnees : ligne " + str(derniere_ligne) + " (sur " + str(ws.max_row) + ")")
        
        # Detecter les colonnes
        headers = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(h).strip().upper() if h else '' for h in row]
            break
        
        idx = {}
        for i, h in enumerate(headers):
            if ('SOUS ACTIVITE' in h or 'SOUS-ACTIVITE' in h) and 'sous_activite' not in idx:
                idx['sous_activite'] = i
            elif 'PROGRAMME' in h and 'programme' not in idx:
                idx['programme'] = i
            elif ('ACTIVITE' in h or 'ACTIVITÉ' in h) and 'activite' not in idx:
                idx['activite'] = i
            elif 'PRODUIT' in h and 'produit' not in idx:
                idx['produit'] = i
            elif ('UNITE' in h or 'UNITÉ' in h) and 'unite' not in idx:
                idx['unite'] = i
            elif 'CIBLE' in h and 'cible' not in idx:
                idx['cible'] = i
            elif 'SOURCE' in h and 'source' not in idx:
                idx['source'] = i
            elif ('PCOP_CODE' in h or ('PCOP' in h and 'CODE' in h)) and 'pcop_code' not in idx:
                idx['pcop_code'] = i
            elif ('PCOP_LIB' in h or ('PCOP' in h and 'LIB' in h)) and 'pcop_lib' not in idx:
                idx['pcop_lib'] = i
            # NOUVELLE DETECTION : Montant et Estimation
            elif 'MONTANT' in h and 'montant' not in idx:
                idx['montant'] = i
            elif 'ESTIMATION' in h and 'montant' not in idx:
                idx['montant'] = i
        
        print("    Colonnes : " + str(idx))
        
        # Memoire pour cellules fusionnees
        dernier_programme = None
        derniere_activite = None
        derniere_sous_activite = None
        
        for row_num, row in enumerate(ws.iter_rows(min_row=2, max_row=derniere_ligne, values_only=True), start=2):
            if not row or len(row) < 4:
                continue
            
            def gv(key):
                i = idx.get(key, -1)
                return row[i] if 0 <= i < len(row) else None
            
            programme_code = gv('programme')
            activite_lib = gv('activite')
            sous_lib = gv('sous_activite')
            produit_lib = gv('produit')
            unite = gv('unite')
            cible = gv('cible')
            source = gv('source')
            pcop_code = gv('pcop_code')
            pcop_lib = gv('pcop_lib')
            montant = gv('montant')  # NOUVEAU
            
            # Heriter des valeurs precedentes (cellules fusionnees)
            if programme_code:
                dernier_programme = programme_code
            else:
                programme_code = dernier_programme
            
            if activite_lib:
                derniere_activite = activite_lib
            else:
                activite_lib = derniere_activite
            
            if sous_lib:
                derniere_sous_activite = sous_lib
            else:
                sous_lib = derniere_sous_activite
            
            # Ignorer lignes vides
            if not any([programme_code, activite_lib, produit_lib, sous_lib]):
                total_ignorees += 1
                continue
            
            # Ignorer les formules
            if isinstance(programme_code, str) and programme_code.startswith('='):
                continue
            
            try:
                # Programme
                try:
                    code_str = str(int(float(programme_code))) if programme_code else None
                except (ValueError, TypeError):
                    code_str = None
                
                if not code_str:
                    total_ignorees += 1
                    continue
                
                programme, _ = Programme.objects.get_or_create(
                    code=code_str,
                    defaults={'libelle': 'Programme ' + code_str}
                )
                
                # Activite
                if not activite_lib:
                    total_ignorees += 1
                    continue
                activite, _ = Activite.objects.get_or_create(
                    programme=programme,
                    libelle=str(activite_lib).strip()[:300]
                )
                
                # CORRECTION : Si produit vide mais sous_lib rempli,
                # alors sous_lib est en fait un PRODUIT
                if not produit_lib and sous_lib:
                    produit_lib = sous_lib
                    sous_lib = None
                
                # Sous-activite : UNE SEULE par defaut
                if sous_lib:
                    sous_act, _ = SousActivite.objects.get_or_create(
                        activite=activite,
                        libelle=str(sous_lib).strip()[:300]
                    )
                else:
                    sous_act, _ = SousActivite.objects.get_or_create(
                        activite=activite,
                        libelle="General"
                    )
                
                # Produit
                if not produit_lib:
                    total_ignorees += 1
                    continue
                
                unite_norm = str(unite).strip() if unite else 'Nombre'
                if unite_norm not in ['Nombre', 'Pack', 'Litre', 'Jour', 'Carte', 'Piece', 'Sac', 'm3']:
                    unite_norm = 'Nombre'
                try:
                    cible_val = Decimal(str(cible)) if cible else Decimal('0')
                except (ValueError, TypeError):
                    cible_val = Decimal('0')
                
                produit, prod_created = Produit.objects.get_or_create(
                    sous_activite=sous_act,
                    libelle=str(produit_lib).strip()[:300],
                    defaults={'unite': unite_norm, 'cible': cible_val}
                )
                if prod_created:
                    total_produits += 1
                
                # PCOP
                pcop = None
                if pcop_code:
                    try:
                        pc_str = str(int(float(pcop_code)))
                    except (ValueError, TypeError):
                        pc_str = str(pcop_code).strip()
                    pcop, _ = PCOP.objects.get_or_create(
                        code=pc_str[:10],
                        defaults={'libelle': str(pcop_lib)[:300] if pcop_lib else 'Code ' + pc_str}
                    )
                
                # NOUVEAU : Montant
                try:
                    montant_val = Decimal(str(montant)) if montant else Decimal('0')
                except (ValueError, TypeError):
                    montant_val = Decimal('0')
                
                # Ligne budgetaire
                source_fin = 'FCE' if (source and 'FCE' in str(source).upper()) else 'RPI'
                ligne, l_created = LigneBudgetaire.objects.get_or_create(
                    etablissement=etab,
                    produit=produit,
                    source_financement=source_fin,
                    defaults={
                        'pcop': pcop,
                        'statut': 'Planifie',
                        'montant_prevu': montant_val,
                    }
                )
                if l_created:
                    total_lignes += 1
                
                # Mise a jour si montant = 0 mais montant_val > 0
                if not l_created and ligne.montant_prevu == 0 and montant_val > 0:
                    ligne.montant_prevu = montant_val
                    ligne.save()
                
                if montant_val > 0:
                    total_montants += 1
            
            except Exception as e:
                print("    ! Erreur ligne " + str(row_num) + " : " + str(e))
                continue
    
    print("    --> " + str(total_produits) + " produits, " + str(total_lignes) + " lignes, " + str(total_montants) + " montants, " + str(total_ignorees) + " ignorees")
    return total_produits, total_lignes, total_montants, total_ignorees


def importer_tous():
    print("=" * 60)
    print("  IMPORTATION DES CANEVAS (VERSION FINALE)")
    print("=" * 60)
    
    if not os.path.exists(DOSSIER_CANEVAS):
        print("ERREUR : Dossier " + DOSSIER_CANEVAS + " introuvable.")
        return
    
    fichiers = sorted(glob.glob(os.path.join(DOSSIER_CANEVAS, '*.xlsx')))
    
    if not fichiers:
        print("ERREUR : Aucun fichier xlsx dans " + DOSSIER_CANEVAS)
        return
    
    print("")
    print(str(len(fichiers)) + " fichier(s) trouve(s)")
    
    total_p = 0
    total_l = 0
    total_m = 0
    total_i = 0
    
    for f in fichiers:
        p, l, m, i = importer_fichier(f)
        total_p += p
        total_l += l
        total_m += m
        total_i += i
    
    print("")
    print("=" * 60)
    print("  RESUME GLOBAL")
    print("=" * 60)
    print("  Fichiers traites     : " + str(len(fichiers)))
    print("  Nouveaux produits    : " + str(total_p))
    print("  Nouvelles lignes     : " + str(total_l))
    print("  Montants renseignes  : " + str(total_m))
    print("  Lignes ignorees      : " + str(total_i))
    print("")
    print("  Etat de la base :")
    print("     Etablissements     : " + str(Etablissement.objects.count()))
    print("     Programmes         : " + str(Programme.objects.count()))
    print("     Activites          : " + str(Activite.objects.count()))
    print("     Sous-activites     : " + str(SousActivite.objects.count()))
    print("     Produits           : " + str(Produit.objects.count()))
    print("     Codes PCOP         : " + str(PCOP.objects.count()))
    print("     Lignes budgetaires : " + str(LigneBudgetaire.objects.count()))
    print("")
    
    # Total des montants
    from django.db.models import Sum
    total_montant = LigneBudgetaire.objects.aggregate(Sum('montant_prevu'))['montant_prevu__sum'] or 0
    print("  MONTANT TOTAL PREVU : " + str(total_montant) + " Ar")
    print("")
    print("IMPORTATION TERMINEE !")


if __name__ == '__main__':
    importer_tous()