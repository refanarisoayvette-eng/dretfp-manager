"""
Diagnostic pour comprendre pourquoi tant de lignes sont ignorées.
"""

import os
import glob
import openpyxl
from collections import Counter


DOSSIER_CANEVAS = 'donnees/canevas'


def diagnostiquer_fichier(chemin):
    """Analyse un fichier et compte les raisons d'ignorance."""
    nom_fichier = os.path.basename(chemin)
    
    print(f"\n{'=' * 70}")
    print(f"  DIAGNOSTIC : {nom_fichier}")
    print(f"{'=' * 70}")
    
    wb = openpyxl.load_workbook(chemin, data_only=True)
    
    for nom_feuille in wb.sheetnames:
        ws = wb[nom_feuille]
        print(f"\n  Feuille : '{nom_feuille}' ({ws.max_row} lignes)")
        print(f"  {'─' * 66}")
        
        # Détecter les colonnes
        headers = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(h).strip().upper() if h else '' for h in row]
            break
        
        idx = {}
        for i, h in enumerate(headers):
            if ('SOUS ACTIVITE' in h or 'SOUS-ACTIVITE' in h) and 'sous_activite' not in idx:
                idx['sous_activite'] = i
            elif 'PROGRAMME' in h and 'programme' not in idx:
                idx['programme'] = i
            elif ('ACTIVITE' in h or 'ACTIVITÉ' in h) and 'activite' not in idx:
                idx['activite'] = i
            elif 'PRODUIT' in h and 'produit' not in idx:
                idx['produit'] = i
            elif ('UNITE' in h or 'UNITÉ' in h) and 'unite' not in idx:
                idx['unite'] = i
            elif 'CIBLE' in h and 'cible' not in idx:
                idx['cible'] = i
            elif 'SOURCE' in h and 'source' not in idx:
                idx['source'] = i
            elif ('PCOP_CODE' in h or ('PCOP' in h and 'CODE' in h)) and 'pcop_code' not in idx:
                idx['pcop_code'] = i
            elif ('PCOP_LIB' in h or ('PCOP' in h and 'LIB' in h)) and 'pcop_lib' not in idx:
                idx['pcop_lib'] = i
        
        print(f"  Colonnes détectées : {idx}")
        
        # Compteurs
        total = 0
        vides = 0
        sans_programme = 0
        sans_activite = 0
        sans_produit = 0
        valides = 0
        lignes_avec_programme = 0
        lignes_avec_activite = 0
        lignes_avec_produit = 0
        
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not any(row):
                continue
            
            total += 1
            
            def gv(key):
                i = idx.get(key, -1)
                return row[i] if 0 <= i < len(row) else None
            
            programme_code = gv('programme')
            activite_lib = gv('activite')
            sous_lib = gv('sous_activite')
            produit_lib = gv('produit')
            
            # Ignorer les formules
            if isinstance(programme_code, str) and programme_code.startswith('='):
                continue
            
            if programme_code:
                lignes_avec_programme += 1
            if activite_lib:
                lignes_avec_activite += 1
            if produit_lib:
                lignes_avec_produit += 1
            
            # Raisons d'ignorance
            if not any([activite_lib, produit_lib, sous_lib]):
                vides += 1
                continue
            
            if not programme_code:
                sans_programme += 1
                continue
            
            if not activite_lib:
                sans_activite += 1
                continue
            
            valides += 1
        
        print(f"\n  STATISTIQUES :")
        print(f"    Total lignes           : {total}")
        print(f"    Lignes AVEC programme  : {lignes_avec_programme}")
        print(f"    Lignes AVEC activité   : {lignes_avec_activite}")
        print(f"    Lignes AVEC produit    : {lignes_avec_produit}")
        print(f"    {'─' * 60}")
        print(f"    ✅ Valides (importables) : {valides}")
        print(f"    ❌ Vides (ignorées)     : {vides}")
        print(f"    ❌ Sans programme       : {sans_programme}")
        print(f"    ❌ Sans activité        : {sans_activite}")


def main():
    print("=" * 70)
    print("  DIAGNOSTIC D'IMPORTATION")
    print("=" * 70)
    
    fichiers = sorted(glob.glob(os.path.join(DOSSIER_CANEVAS, '*.xlsx')))
    
    for f in fichiers:
        diagnostiquer_fichier(f)


if __name__ == '__main__':
    main()