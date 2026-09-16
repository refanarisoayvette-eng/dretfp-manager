from django.shortcuts import render, get_object_or_404, redirect
from django.http import JsonResponse, HttpResponse, FileResponse
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import UserCreationForm
from django.contrib.auth import login
from django.contrib.auth.models import User
from django import forms
from django.db.models import Sum, Count, F, Q
from django.core.paginator import Paginator
from .models import (
    Etablissement, Produit, LigneBudgetaire, Activite,
    Programme, SousActivite, PCOP, Profil, DemandeReinitialisation,
    PTA, ActivitePTA,
    Salle, Filiere, Matiere, CreneauHoraire, EmploiDuTemps
)
from .decorators import role_required

# Imports PDF
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_LEFT
import io
from datetime import datetime


# ============================================================
# PAGES PRINCIPALES
# ============================================================

@login_required
def accueil(request):
    context = {
        'nb_etablissements': Etablissement.objects.count(),
        'nb_produits': Produit.objects.count(),
        'nb_lignes': LigneBudgetaire.objects.count(),
        'nb_activites': Activite.objects.count(),
    }
    return render(request, 'core/accueil.html', context)


@login_required
def tableau_bord(request):
    """Tableau de bord complet avec KPI, graphiques et filtres."""
    etab_id = request.GET.get('etablissement')
    source = request.GET.get('source')
    
    lignes = LigneBudgetaire.objects.all()
    if etab_id:
        lignes = lignes.filter(etablissement_id=etab_id)
    if source:
        lignes = lignes.filter(source_financement=source)
    
    # KPI
    total_prevu = lignes.aggregate(Sum('montant_prevu'))['montant_prevu__sum'] or 0
    total_realise = lignes.aggregate(Sum('montant_realise'))['montant_realise__sum'] or 0
    nb_lignes = lignes.count()
    nb_lignes_avec_montant = lignes.filter(montant_prevu__gt=0).count()
    nb_lignes_sans_montant = nb_lignes - nb_lignes_avec_montant
    
    # KPI globaux
    total_etab = Etablissement.objects.count()
    total_produits = Produit.objects.count()
    total_activites = Activite.objects.count()
    
    # RPI vs FCE
    rpi_total = LigneBudgetaire.objects.filter(source_financement='RPI').aggregate(Sum('montant_prevu'))['montant_prevu__sum'] or 0
    fce_total = LigneBudgetaire.objects.filter(source_financement='FCE').aggregate(Sum('montant_prevu'))['montant_prevu__sum'] or 0
    
    taux = 0
    if total_prevu > 0:
        taux = round((float(total_realise) / float(total_prevu)) * 100, 2)
    
    # Alertes
    alertes = []
    if nb_lignes_sans_montant > 0:
        alertes.append({'type': 'warning', 'message': str(nb_lignes_sans_montant) + " ligne(s) sans montant prevu"})
    
    depassements = lignes.filter(montant_realise__gt=F('montant_prevu')).count()
    if depassements > 0:
        alertes.append({'type': 'danger', 'message': str(depassements) + " ligne(s) en depassement budgetaire"})
    
    # Stats
    stats_statuts = lignes.values('statut').annotate(count=Count('id'), total=Sum('montant_prevu')).order_by('statut')
    top_etabs = LigneBudgetaire.objects.values('etablissement__nom').annotate(
        total=Sum('montant_prevu'), nb_lignes=Count('id')
    ).order_by('-total')[:5]
    
    context = {
        'total_prevu': total_prevu, 'total_realise': total_realise,
        'taux_execution': taux, 'nb_lignes': nb_lignes,
        'nb_lignes_avec_montant': nb_lignes_avec_montant,
        'nb_lignes_sans_montant': nb_lignes_sans_montant,
        'total_etab': total_etab, 'total_produits': total_produits,
        'total_activites': total_activites,
        'rpi_total': rpi_total, 'fce_total': fce_total,
        'alertes': alertes, 'etablissements': Etablissement.objects.all(),
        'stats_statuts': list(stats_statuts), 'top_etabs': list(top_etabs),
        'filtre_etab': etab_id or '', 'filtre_source': source or '',
    }
    return render(request, 'core/tableau_bord.html', context)


@login_required
def statistiques(request):
    """Redirection vers le tableau de bord (fusionne)."""
    return redirect('tableau_bord')


@login_required
def liste_etablissements(request):
    etablissements = Etablissement.objects.all()
    for etab in etablissements:
        etab.nb_produits = Produit.objects.filter(
            sous_activite__activite__programme__code='319'
        ).filter(lignebudgetaire__etablissement=etab).distinct().count()
        etab.nb_lignes = LigneBudgetaire.objects.filter(etablissement=etab).count()
    return render(request, 'core/liste_etablissements.html', {'etablissements': etablissements})


# ============================================================
# BUDGETS (CRUD)
# ============================================================

@login_required
def liste_budgets(request):
    lignes = LigneBudgetaire.objects.select_related('etablissement', 'produit', 'pcop').all()
    etab_id = request.GET.get('etablissement')
    if etab_id:
        lignes = lignes.filter(etablissement_id=etab_id)
    source = request.GET.get('source')
    if source:
        lignes = lignes.filter(source_financement=source)
    statut = request.GET.get('statut')
    if statut:
        lignes = lignes.filter(statut=statut)
    q = request.GET.get('q')
    if q:
        lignes = lignes.filter(Q(produit__libelle__icontains=q) | Q(etablissement__nom__icontains=q) | Q(remarque__icontains=q))
    montant_min = request.GET.get('montant_min')
    if montant_min:
        try:
            lignes = lignes.filter(montant_prevu__gte=float(montant_min))
        except (ValueError, TypeError):
            pass
    tri = request.GET.get('tri', '-date_creation')
    tris_valides = ['produit__libelle', '-produit__libelle', 'montant_prevu', '-montant_prevu',
                    'montant_realise', '-montant_realise', 'date_creation', '-date_creation',
                    'etablissement__nom', '-etablissement__nom']
    if tri in tris_valides:
        lignes = lignes.order_by(tri)
    paginator = Paginator(lignes, 25)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    context = {
        'lignes': page_obj, 'page_obj': page_obj, 'paginator': paginator,
        'etablissements': Etablissement.objects.all(),
        'filtre_q': q or '', 'filtre_etab': etab_id or '',
        'filtre_source': source or '', 'filtre_statut': statut or '',
        'filtre_montant_min': montant_min or '', 'filtre_tri': tri,
    }
    return render(request, 'core/liste_budgets.html', context)


@login_required
def editer_budget(request, ligne_id):
    ligne = get_object_or_404(LigneBudgetaire, id=ligne_id)
    if request.method == 'POST':
        try:
            ligne.montant_prevu = request.POST.get('montant_prevu', 0) or 0
            ligne.montant_realise = request.POST.get('montant_realise', 0) or 0
            ligne.statut = request.POST.get('statut', 'Planifie')
            ligne.remarque = request.POST.get('remarque', '')
            ligne.save()
            messages.success(request, "Ligne budgetaire modifiee !")
        except Exception as e:
            messages.error(request, "Erreur : " + str(e))
        return redirect('liste_budgets')
    return render(request, 'core/editer_budget.html', {'ligne': ligne})


@login_required
def ajouter_budget(request):
    if request.method == 'POST':
        etab_id = request.POST.get('etablissement')
        produit_id = request.POST.get('produit')
        if not etab_id or not produit_id:
            messages.error(request, "Veuillez remplir tous les champs obligatoires.")
        else:
            try:
                ligne = LigneBudgetaire(
                    etablissement_id=etab_id, produit_id=produit_id,
                    source_financement=request.POST.get('source_financement', 'RPI'),
                    montant_prevu=request.POST.get('montant_prevu', 0) or 0,
                    montant_realise=request.POST.get('montant_realise', 0) or 0,
                    statut=request.POST.get('statut', 'Planifie'),
                    remarque=request.POST.get('remarque', ''),
                )
                pcop_id = request.POST.get('pcop')
                if pcop_id:
                    ligne.pcop_id = pcop_id
                ligne.save()
                messages.success(request, "Ligne budgetaire ajoutee !")
                return redirect('liste_budgets')
            except Exception as e:
                messages.error(request, "Erreur : " + str(e))
    context = {
        'etablissements': Etablissement.objects.all(),
        'produits': Produit.objects.all()[:500],
        'pcops': PCOP.objects.all(),
    }
    return render(request, 'core/ajouter_budget.html', context)


@login_required
def supprimer_budget(request, ligne_id):
    ligne = get_object_or_404(LigneBudgetaire, id=ligne_id)
    if request.method == 'POST':
        etab = ligne.etablissement.nom
        produit = ligne.produit.libelle
        ligne.delete()
        messages.success(request, "Ligne supprimee : " + etab + " - " + produit)
        return redirect('liste_budgets')
    return render(request, 'core/supprimer_budget.html', {'ligne': ligne})


# ============================================================
# ACTIVITES (CRUD)
# ============================================================

@login_required
def liste_activites(request):
    activites = Activite.objects.select_related('programme').all()
    prog_id = request.GET.get('programme')
    if prog_id:
        activites = activites.filter(programme_id=prog_id)
    q = request.GET.get('q')
    if q:
        activites = activites.filter(libelle__icontains=q)
    context = {
        'activites': activites, 'programmes': Programme.objects.all(),
        'filtre_prog': prog_id or '', 'filtre_q': q or '',
    }
    return render(request, 'core/liste_activites.html', context)


@login_required
@role_required('DRETFP')
def ajouter_activite(request):
    if request.method == 'POST':
        libelle = request.POST.get('libelle', '').strip()
        programme_id = request.POST.get('programme')
        objectif = request.POST.get('objectif', '').strip()
        if not libelle or not programme_id:
            messages.error(request, "Le libelle et le programme sont obligatoires.")
        else:
            try:
                Activite.objects.create(libelle=libelle, programme_id=programme_id, objectif=objectif)
                messages.success(request, "Activite ajoutee : " + libelle)
                return redirect('liste_activites')
            except Exception as e:
                messages.error(request, "Erreur : " + str(e))
    return render(request, 'core/ajouter_activite.html', {'programmes': Programme.objects.all()})


@login_required
@role_required('DRETFP')
def editer_activite(request, activite_id):
    activite = get_object_or_404(Activite, id=activite_id)
    if request.method == 'POST':
        activite.libelle = request.POST.get('libelle', activite.libelle)
        activite.objectif = request.POST.get('objectif', '')
        programme_id = request.POST.get('programme')
        if programme_id:
            activite.programme_id = programme_id
        activite.save()
        messages.success(request, "Activite modifiee !")
        return redirect('liste_activites')
    return render(request, 'core/editer_activite.html', {'activite': activite, 'programmes': Programme.objects.all()})


@login_required
@role_required('DRETFP')
def supprimer_activite(request, activite_id):
    activite = get_object_or_404(Activite, id=activite_id)
    if request.method == 'POST':
        libelle = activite.libelle
        activite.delete()
        messages.success(request, "Activite supprimee : " + libelle)
        return redirect('liste_activites')
    return render(request, 'core/supprimer_activite.html', {'activite': activite})


# ============================================================
# PTA
# ============================================================

@login_required
def liste_pta(request):
    ptas = PTA.objects.select_related('etablissement').all()
    etab_id = request.GET.get('etablissement')
    if etab_id:
        ptas = ptas.filter(etablissement_id=etab_id)
    annee = request.GET.get('annee')
    if annee:
        ptas = ptas.filter(annee=annee)
    statut = request.GET.get('statut')
    if statut:
        ptas = ptas.filter(statut=statut)
    if not request.user.is_staff:
        try:
            profil = request.user.profil
            if profil.role == 'CHEF_ETAB' and profil.etablissement:
                ptas = ptas.filter(etablissement=profil.etablissement)
        except:
            pass
    context = {
        'ptas': ptas, 'etablissements': Etablissement.objects.all(),
        'filtre_etab': etab_id or '', 'filtre_annee': annee or '', 'filtre_statut': statut or '',
        'annees': [2025, 2026, 2027, 2028, 2029, 2030],
    }
    return render(request, 'core/liste_pta.html', context)


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def ajouter_pta(request):
    if request.method == 'POST':
        etablissement_id = request.POST.get('etablissement')
        annee = request.POST.get('annee')
        titre = request.POST.get('titre', '').strip()
        if not etablissement_id or not annee or not titre:
            messages.error(request, "L'etablissement, l'annee et le titre sont obligatoires.")
        else:
            try:
                pta = PTA.objects.create(
                    etablissement_id=etablissement_id, annee=annee, titre=titre,
                    responsable=request.POST.get('responsable', '').strip(),
                    date_debut=request.POST.get('date_debut') or None,
                    date_fin=request.POST.get('date_fin') or None,
                    remarque=request.POST.get('remarque', '').strip(),
                    statut='Brouillon',
                )
                messages.success(request, "PTA cree : " + titre)
                return redirect('detail_pta', pta_id=pta.id)
            except Exception as e:
                messages.error(request, "Erreur : " + str(e))
    context = {'etablissements': Etablissement.objects.all(), 'annees': [2025, 2026, 2027, 2028, 2029, 2030]}
    return render(request, 'core/ajouter_pta.html', context)


@login_required
def detail_pta(request, pta_id):
    pta = get_object_or_404(PTA, id=pta_id)
    activites = pta.activites.all().order_by('ordre', 'date_debut')
    return render(request, 'core/detail_pta.html', {'pta': pta, 'activites': activites})


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def editer_pta(request, pta_id):
    pta = get_object_or_404(PTA, id=pta_id)
    if request.method == 'POST':
        pta.titre = request.POST.get('titre', pta.titre)
        pta.responsable = request.POST.get('responsable', '')
        pta.statut = request.POST.get('statut', pta.statut)
        pta.date_debut = request.POST.get('date_debut') or None
        pta.date_fin = request.POST.get('date_fin') or None
        pta.remarque = request.POST.get('remarque', '')
        pta.save()
        messages.success(request, "PTA modifie !")
        return redirect('detail_pta', pta_id=pta.id)
    context = {'pta': pta, 'etablissements': Etablissement.objects.all(), 'annees': [2025, 2026, 2027, 2028, 2029, 2030]}
    return render(request, 'core/editer_pta.html', context)


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def supprimer_pta(request, pta_id):
    pta = get_object_or_404(PTA, id=pta_id)
    if request.method == 'POST':
        titre = pta.titre
        pta.delete()
        messages.success(request, "PTA supprime : " + titre)
        return redirect('liste_pta')
    return render(request, 'core/supprimer_pta.html', {'pta': pta})


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def ajouter_activite_pta(request, pta_id):
    pta = get_object_or_404(PTA, id=pta_id)
    if request.method == 'POST':
        titre = request.POST.get('titre', '').strip()
        if not titre:
            messages.error(request, "Le titre est obligatoire.")
        else:
            try:
                ActivitePTA.objects.create(
                    PTA=pta, titre=titre,
                    description=request.POST.get('description', '').strip(),
                    responsable=request.POST.get('responsable', '').strip(),
                    statut=request.POST.get('statut', 'Non commence'),
                    taux_avancement=int(request.POST.get('taux_avancement', 0) or 0),
                    date_debut=request.POST.get('date_debut') or None,
                    date_fin=request.POST.get('date_fin') or None,
                    budget_prevu=request.POST.get('budget_prevu', 0) or 0,
                    budget_realise=request.POST.get('budget_realise', 0) or 0,
                    ordre=int(request.POST.get('ordre', 0) or 0),
                )
                messages.success(request, "Activite ajoutee : " + titre)
                return redirect('detail_pta', pta_id=pta.id)
            except Exception as e:
                messages.error(request, "Erreur : " + str(e))
    return render(request, 'core/ajouter_activite_pta.html', {'pta': pta})


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def editer_activite_pta(request, activite_id):
    activite = get_object_or_404(ActivitePTA, id=activite_id)
    if request.method == 'POST':
        activite.titre = request.POST.get('titre', activite.titre)
        activite.description = request.POST.get('description', '')
        activite.responsable = request.POST.get('responsable', '')
        activite.statut = request.POST.get('statut', activite.statut)
        activite.taux_avancement = int(request.POST.get('taux_avancement', 0) or 0)
        activite.date_debut = request.POST.get('date_debut') or None
        activite.date_fin = request.POST.get('date_fin') or None
        activite.budget_prevu = request.POST.get('budget_prevu', 0) or 0
        activite.budget_realise = request.POST.get('budget_realise', 0) or 0
        activite.ordre = int(request.POST.get('ordre', 0) or 0)
        activite.save()
        messages.success(request, "Activite modifiee !")
        return redirect('detail_pta', pta_id=activite.PTA.id)
    return render(request, 'core/editer_activite_pta.html', {'activite': activite})


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def supprimer_activite_pta(request, activite_id):
    activite = get_object_or_404(ActivitePTA, id=activite_id)
    pta_id = activite.PTA.id
    if request.method == 'POST':
        titre = activite.titre
        activite.delete()
        messages.success(request, "Activite supprimee : " + titre)
        return redirect('detail_pta', pta_id=pta_id)
    return render(request, 'core/supprimer_activite_pta.html', {'activite': activite})


# ============================================================
# SALLES (CRUD)
# ============================================================

@login_required
def liste_salles(request):
    salles = Salle.objects.select_related('etablissement').all()
    etab_id = request.GET.get('etablissement')
    if etab_id:
        salles = salles.filter(etablissement_id=etab_id)
    context = {
        'salles': salles, 'etablissements': Etablissement.objects.all(),
        'filtre_etab': etab_id or '',
    }
    return render(request, 'core/liste_salles.html', context)


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def ajouter_salle(request):
    if request.method == 'POST':
        try:
            Salle.objects.create(
                etablissement_id=request.POST.get('etablissement'),
                nom=request.POST.get('nom'),
                capacite=request.POST.get('capacite', 30) or 30,
                type_salle=request.POST.get('type_salle', ''),
                batiment=request.POST.get('batiment', ''),
            )
            messages.success(request, "Salle ajoutee !")
            return redirect('liste_salles')
        except Exception as e:
            messages.error(request, "Erreur : " + str(e))
    return render(request, 'core/ajouter_salle.html', {'etablissements': Etablissement.objects.all()})


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def editer_salle(request, salle_id):
    salle = get_object_or_404(Salle, id=salle_id)
    if request.method == 'POST':
        salle.nom = request.POST.get('nom', salle.nom)
        salle.capacite = request.POST.get('capacite', salle.capacite) or 30
        salle.type_salle = request.POST.get('type_salle', '')
        salle.batiment = request.POST.get('batiment', '')
        salle.actif = 'actif' in request.POST
        salle.save()
        messages.success(request, "Salle modifiee !")
        return redirect('liste_salles')
    return render(request, 'core/editer_salle.html', {'salle': salle})


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def supprimer_salle(request, salle_id):
    salle = get_object_or_404(Salle, id=salle_id)
    if request.method == 'POST':
        salle.delete()
        messages.success(request, "Salle supprimee !")
        return redirect('liste_salles')
    return render(request, 'core/supprimer_salle.html', {'salle': salle})


# ============================================================
# FILIERES (CRUD)
# ============================================================

@login_required
def liste_filieres(request):
    filieres = Filiere.objects.select_related('etablissement').all()
    etab_id = request.GET.get('etablissement')
    if etab_id:
        filieres = filieres.filter(etablissement_id=etab_id)
    context = {
        'filieres': filieres, 'etablissements': Etablissement.objects.all(),
        'filtre_etab': etab_id or '',
    }
    return render(request, 'core/liste_filieres.html', context)


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def ajouter_filiere(request):
    if request.method == 'POST':
        try:
            Filiere.objects.create(
                etablissement_id=request.POST.get('etablissement'),
                code=request.POST.get('code'),
                nom=request.POST.get('nom'),
                niveau=request.POST.get('niveau', 'CAP'),
                nb_apprenants=request.POST.get('nb_apprenants', 0) or 0,
                duree_annees=request.POST.get('duree_annees', 2) or 2,
            )
            messages.success(request, "Filiere ajoutee !")
            return redirect('liste_filieres')
        except Exception as e:
            messages.error(request, "Erreur : " + str(e))
    return render(request, 'core/ajouter_filiere.html', {'etablissements': Etablissement.objects.all()})


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def editer_filiere(request, filiere_id):
    filiere = get_object_or_404(Filiere, id=filiere_id)
    if request.method == 'POST':
        filiere.code = request.POST.get('code', filiere.code)
        filiere.nom = request.POST.get('nom', filiere.nom)
        filiere.niveau = request.POST.get('niveau', filiere.niveau)
        filiere.nb_apprenants = request.POST.get('nb_apprenants', 0) or 0
        filiere.duree_annees = request.POST.get('duree_annees', 2) or 2
        filiere.actif = 'actif' in request.POST
        filiere.save()
        messages.success(request, "Filiere modifiee !")
        return redirect('liste_filieres')
    return render(request, 'core/editer_filiere.html', {'filiere': filiere})


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def supprimer_filiere(request, filiere_id):
    filiere = get_object_or_404(Filiere, id=filiere_id)
    if request.method == 'POST':
        filiere.delete()
        messages.success(request, "Filiere supprimee !")
        return redirect('liste_filieres')
    return render(request, 'core/supprimer_filiere.html', {'filiere': filiere})


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def creer_filiere_ajax(request):
    """Creer une filiere en AJAX depuis la page matiere."""
    if request.method == 'POST':
        try:
            etablissement_id = request.POST.get('etablissement')
            code = request.POST.get('code', '').strip()
            nom = request.POST.get('nom', '').strip()
            niveau = request.POST.get('niveau', 'CAP')
            nb_apprenants = request.POST.get('nb_apprenants', 0) or 0
            duree_annees = request.POST.get('duree_annees', 2) or 2
            
            if not etablissement_id or not code or not nom:
                return JsonResponse({'success': False, 'error': 'Tous les champs sont obligatoires'})
            
            filiere = Filiere.objects.create(
                etablissement_id=etablissement_id, code=code, nom=nom,
                niveau=niveau, nb_apprenants=int(nb_apprenants), duree_annees=int(duree_annees),
            )
            
            return JsonResponse({
                'success': True, 'id': filiere.id, 'code': filiere.code, 'nom': filiere.nom,
                'etablissement': filiere.etablissement.nom,
                'label': f"{filiere.etablissement.nom} - {filiere.code} - {filiere.nom}"
            })
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)})
    return JsonResponse({'success': False, 'error': 'Methode non autorisee'})


# ============================================================
# MATIERES (CRUD)
# ============================================================

@login_required
def liste_matieres(request):
    matieres = Matiere.objects.select_related('filiere', 'filiere__etablissement').all()
    filiere_id = request.GET.get('filiere')
    if filiere_id:
        matieres = matieres.filter(filiere_id=filiere_id)
    etab_id = request.GET.get('etablissement')
    if etab_id:
        matieres = matieres.filter(filiere__etablissement_id=etab_id)
    context = {
        'matieres': matieres, 'filieres': Filiere.objects.all(),
        'etablissements': Etablissement.objects.all(),
        'filtre_filiere': filiere_id or '', 'filtre_etab': etab_id or '',
    }
    return render(request, 'core/liste_matieres.html', context)


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def ajouter_matiere(request):
    if request.method == 'POST':
        try:
            Matiere.objects.create(
                filiere_id=request.POST.get('filiere'),
                code=request.POST.get('code'),
                nom=request.POST.get('nom'),
                type_matiere=request.POST.get('type_matiere', 'Theorique'),
                heures_theorie=request.POST.get('heures_theorie', 0) or 0,
                heures_pratique=request.POST.get('heures_pratique', 0) or 0,
                coefficient=request.POST.get('coefficient', 1) or 1,
                description=request.POST.get('description', ''),
            )
            messages.success(request, "Matiere ajoutee !")
            return redirect('liste_matieres')
        except Exception as e:
            messages.error(request, "Erreur : " + str(e))
    filiere_id = request.GET.get('filiere')
    context = {
        'filieres': Filiere.objects.all(),
        'filiere_selectionnee': filiere_id,
        'etablissements_modal': Etablissement.objects.all(),
    }
    return render(request, 'core/ajouter_matiere.html', context)


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def editer_matiere(request, matiere_id):
    matiere = get_object_or_404(Matiere, id=matiere_id)
    if request.method == 'POST':
        matiere.filiere_id = request.POST.get('filiere', matiere.filiere_id)
        matiere.code = request.POST.get('code', matiere.code)
        matiere.nom = request.POST.get('nom', matiere.nom)
        matiere.type_matiere = request.POST.get('type_matiere', matiere.type_matiere)
        matiere.heures_theorie = request.POST.get('heures_theorie', 0) or 0
        matiere.heures_pratique = request.POST.get('heures_pratique', 0) or 0
        matiere.coefficient = request.POST.get('coefficient', 1) or 1
        matiere.description = request.POST.get('description', '')
        matiere.actif = 'actif' in request.POST
        matiere.save()
        messages.success(request, "Matiere modifiee !")
        return redirect('liste_matieres')
    return render(request, 'core/editer_matiere.html', {'matiere': matiere, 'filieres': Filiere.objects.all()})


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def supprimer_matiere(request, matiere_id):
    matiere = get_object_or_404(Matiere, id=matiere_id)
    if request.method == 'POST':
        matiere.delete()
        messages.success(request, "Matiere supprimee !")
        return redirect('liste_matieres')
    return render(request, 'core/supprimer_matiere.html', {'matiere': matiere})


# ============================================================
# FORMATEURS
# ============================================================

@login_required
def liste_formateurs(request):
    formateurs = Profil.objects.filter(role='FORMATEUR').select_related('utilisateur', 'etablissement')
    etab_id = request.GET.get('etablissement')
    if etab_id:
        formateurs = formateurs.filter(etablissement_id=etab_id)
    context = {
        'formateurs': formateurs, 'etablissements': Etablissement.objects.all(),
        'filtre_etab': etab_id or '',
    }
    return render(request, 'core/liste_formateurs.html', context)


# ============================================================
# EMPLOI DU TEMPS
# ============================================================

@login_required
def liste_emplois(request):
    emplois = EmploiDuTemps.objects.select_related(
        'filiere', 'filiere__etablissement', 'salle', 'creneau', 'formateur'
    ).all()
    etab_id = request.GET.get('etablissement')
    filiere_id = request.GET.get('filiere')
    jour = request.GET.get('jour')
    if etab_id:
        emplois = emplois.filter(filiere__etablissement_id=etab_id)
    if filiere_id:
        emplois = emplois.filter(filiere_id=filiere_id)
    if jour:
        emplois = emplois.filter(jour=jour)
    context = {
        'emplois': emplois, 'etablissements': Etablissement.objects.all(),
        'filieres': Filiere.objects.all(),
        'jours': ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi'],
        'filtre_etab': etab_id or '', 'filtre_filiere': filiere_id or '', 'filtre_jour': jour or '',
    }
    return render(request, 'core/liste_emplois.html', context)


@login_required
def vue_hebdomadaire(request):
    filiere_id = request.GET.get('filiere')
    if not filiere_id:
        filiere = Filiere.objects.first()
        if filiere:
            filiere_id = filiere.id
    filiere = get_object_or_404(Filiere, id=filiere_id) if filiere_id else None
    creneaux = CreneauHoraire.objects.all().order_by('ordre')
    jours = ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi']
    grille = {}
    if filiere:
        emplois = EmploiDuTemps.objects.filter(filiere=filiere).select_related('salle', 'creneau', 'formateur')
        for e in emplois:
            grille[(e.jour, e.creneau.id)] = e
    context = {
        'filiere': filiere, 'filieres': Filiere.objects.all(),
        'creneaux': creneaux, 'jours': jours, 'grille': grille,
    }
    return render(request, 'core/vue_hebdomadaire.html', context)


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def ajouter_emploi(request):
    if request.method == 'POST':
        try:
            EmploiDuTemps.objects.create(
                filiere_id=request.POST.get('filiere'),
                creneau_id=request.POST.get('creneau'),
                salle_id=request.POST.get('salle') or None,
                formateur_id=request.POST.get('formateur') or None,
                jour=request.POST.get('jour'),
                matiere=request.POST.get('matiere'),
                annee_scolaire=request.POST.get('annee_scolaire', '2026-2027'),
                remarque=request.POST.get('remarque', ''),
            )
            messages.success(request, "Creneau ajoute !")
            return redirect('liste_emplois')
        except Exception as e:
            messages.error(request, "Erreur : " + str(e))
    formateurs = User.objects.filter(profil__role='FORMATEUR') | User.objects.filter(is_staff=True)
    context = {
        'filieres': Filiere.objects.all(),
        'creneaux': CreneauHoraire.objects.all().order_by('ordre'),
        'salles': Salle.objects.all(), 'formateurs': formateurs,
        'jours': ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi'],
    }
    return render(request, 'core/ajouter_emploi.html', context)


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def editer_emploi(request, emploi_id):
    emploi = get_object_or_404(EmploiDuTemps, id=emploi_id)
    if request.method == 'POST':
        emploi.filiere_id = request.POST.get('filiere', emploi.filiere_id)
        emploi.creneau_id = request.POST.get('creneau', emploi.creneau_id)
        emploi.salle_id = request.POST.get('salle') or None
        emploi.formateur_id = request.POST.get('formateur') or None
        emploi.jour = request.POST.get('jour', emploi.jour)
        emploi.matiere = request.POST.get('matiere', emploi.matiere)
        emploi.annee_scolaire = request.POST.get('annee_scolaire', emploi.annee_scolaire)
        emploi.remarque = request.POST.get('remarque', '')
        emploi.save()
        messages.success(request, "Creneau modifie !")
        return redirect('liste_emplois')
    formateurs = User.objects.filter(profil__role='FORMATEUR') | User.objects.filter(is_staff=True)
    context = {
        'emploi': emploi, 'filieres': Filiere.objects.all(),
        'creneaux': CreneauHoraire.objects.all().order_by('ordre'),
        'salles': Salle.objects.all(), 'formateurs': formateurs,
        'jours': ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi'],
    }
    return render(request, 'core/editer_emploi.html', context)


@login_required
@role_required('DRETFP', 'CHEF_ETAB')
def supprimer_emploi(request, emploi_id):
    emploi = get_object_or_404(EmploiDuTemps, id=emploi_id)
    if request.method == 'POST':
        emploi.delete()
        messages.success(request, "Creneau supprime !")
        return redirect('liste_emplois')
    return render(request, 'core/supprimer_emploi.html', {'emploi': emploi})


# ============================================================
# EXPORTS PTA
# ============================================================

@login_required
def export_pta_pdf(request, pta_id):
    pta = get_object_or_404(PTA, id=pta_id)
    activites = pta.activites.all().order_by('ordre', 'date_debut')
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1*cm, leftMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    elements = []
    styles = getSampleStyleSheet()
    titre_style = ParagraphStyle('Titre', parent=styles['Heading1'], alignment=TA_CENTER,
        textColor=colors.HexColor('#0a2540'), fontSize=16, spaceAfter=10)
    elements.append(Paragraph("DRETFP Manager - Amoron'i Mania", titre_style))
    elements.append(Paragraph("Plan de Travail Annuel (PTA)", styles['Heading2']))
    elements.append(Spacer(1, 0.5*cm))
    infos_data = [
        ['Etablissement', pta.etablissement.nom], ['Annee', str(pta.annee)],
        ['Titre', pta.titre], ['Statut', pta.statut],
        ['Responsable', pta.responsable or '-'],
        ['Avancement global', str(pta.taux_avancement) + ' %'],
        ['Budget total', f"{float(pta.budget_total):,.0f} Ar"],
        ["Nombre d'activites", str(activites.count())],
    ]
    infos_table = Table(infos_data, colWidths=[5*cm, 11*cm])
    infos_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#e7f1ff')),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'), ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('PADDING', (0, 0), (-1, -1), 5), ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    elements.append(infos_table)
    elements.append(Spacer(1, 0.5*cm))
    if activites:
        data = [['#', 'Titre', 'Responsable', 'Debut', 'Fin', 'Statut', 'Avanc.', 'Budget (Ar)']]
        for i, a in enumerate(activites, start=1):
            data.append([str(a.ordre or i), a.titre[:40], (a.responsable or '-')[:15],
                a.date_debut.strftime('%d/%m/%y') if a.date_debut else '-',
                a.date_fin.strftime('%d/%m/%y') if a.date_fin else '-',
                a.statut[:12], str(a.taux_avancement) + '%', f"{float(a.budget_prevu):,.0f}"])
        table = Table(data, colWidths=[0.8*cm, 4*cm, 2*cm, 1.5*cm, 1.5*cm, 2*cm, 1.2*cm, 3*cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a2540')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 7),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
        ]))
        elements.append(table)
    elements.append(Spacer(1, 0.5*cm))
    elements.append(Paragraph(f"Genere le {datetime.now().strftime('%d/%m/%Y a %H:%M')}", styles['Italic']))
    doc.build(elements)
    buffer.seek(0)
    nom_fichier = f"PTA_{pta.etablissement.nom}_{pta.annee}.pdf".replace(' ', '_')
    return FileResponse(buffer, as_attachment=True, filename=nom_fichier)


@login_required
def export_pta_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "PTA DRETFP"
    headers = ['Etablissement', 'Annee', 'Titre', 'Statut', 'Responsable',
               'Date debut', 'Date fin', 'Nb activites', 'Avancement (%)', 'Budget total (Ar)']
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="0a2540", end_color="0a2540", fill_type="solid")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row, pta in enumerate(PTA.objects.select_related('etablissement').all(), start=2):
        ws.cell(row=row, column=1, value=pta.etablissement.nom)
        ws.cell(row=row, column=2, value=pta.annee)
        ws.cell(row=row, column=3, value=pta.titre)
        ws.cell(row=row, column=4, value=pta.statut)
        ws.cell(row=row, column=5, value=pta.responsable or '')
        ws.cell(row=row, column=8, value=pta.nb_activites)
        ws.cell(row=row, column=9, value=pta.taux_avancement)
        ws.cell(row=row, column=10, value=float(pta.budget_total))
    for i, w in enumerate([25, 8, 35, 12, 20, 12, 12, 12, 14, 18], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="PTA_DRETFP.xlsx"'
    wb.save(response)
    return response


@login_required
def export_pta_csv(request):
    import csv
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="PTA_DRETFP.csv"'
    response.write('\ufeff')
    writer = csv.writer(response, delimiter=';')
    writer.writerow(['Etablissement', 'Annee', 'Titre', 'Statut', 'Responsable', 'Nb activites', 'Avancement', 'Budget'])
    for pta in PTA.objects.select_related('etablissement').all():
        writer.writerow([pta.etablissement.nom, pta.annee, pta.titre, pta.statut,
                         pta.responsable or '', pta.nb_activites, pta.taux_avancement, float(pta.budget_total)])
    return response


@login_required
def export_activites_pta_excel(request, pta_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    pta = get_object_or_404(PTA, id=pta_id)
    activites = pta.activites.all().order_by('ordre', 'date_debut')
    wb = Workbook()
    ws = wb.active
    ws.title = f"Activites PTA {pta.annee}"
    headers = ['Ordre', 'Titre', 'Description', 'Responsable', 'Statut',
               'Avancement (%)', 'Date debut', 'Date fin', 'Budget prevu (Ar)', 'Budget realise (Ar)']
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    for row, a in enumerate(activites, start=2):
        ws.cell(row=row, column=1, value=a.ordre)
        ws.cell(row=row, column=2, value=a.titre)
        ws.cell(row=row, column=3, value=a.description)
        ws.cell(row=row, column=4, value=a.responsable)
        ws.cell(row=row, column=5, value=a.statut)
        ws.cell(row=row, column=6, value=a.taux_avancement)
        ws.cell(row=row, column=9, value=float(a.budget_prevu))
        ws.cell(row=row, column=10, value=float(a.budget_realise))
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    nom_fichier = f"Activites_PTA_{pta.etablissement.nom}_{pta.annee}.xlsx".replace(' ', '_')
    response['Content-Disposition'] = f'attachment; filename="{nom_fichier}"'
    wb.save(response)
    return response


# ============================================================
# EXPORTS FILIERES / SALLES / MATIERES / FORMATEURS
# ============================================================

@login_required
def export_filieres_pdf(request):
    filieres = Filiere.objects.select_related('etablissement').all()
    etab_id = request.GET.get('etablissement')
    if etab_id:
        filieres = filieres.filter(etablissement_id=etab_id)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1*cm, leftMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    elements = []
    styles = getSampleStyleSheet()
    titre_style = ParagraphStyle('Titre', parent=styles['Heading1'], alignment=TA_CENTER,
        textColor=colors.HexColor('#0a2540'), fontSize=16, spaceAfter=10)
    elements.append(Paragraph("DRETFP Manager - Amoron'i Mania", titre_style))
    elements.append(Paragraph("Liste des filieres", styles['Heading2']))
    elements.append(Spacer(1, 0.5*cm))
    data = [['Code', 'Nom', 'Etablissement', 'Niveau', 'Apprenants', 'Duree']]
    for f in filieres:
        data.append([f.code, f.nom[:40], f.etablissement.nom[:25], f.get_niveau_display(),
                     str(f.nb_apprenants), f"{f.duree_annees} ans"])
    table = Table(data, colWidths=[2*cm, 5*cm, 4*cm, 2*cm, 2*cm, 1.5*cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a2540')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.5*cm))
    elements.append(Paragraph(f"Total : {filieres.count()} filiere(s)", styles['Normal']))
    elements.append(Paragraph(f"Genere le {datetime.now().strftime('%d/%m/%Y a %H:%M')}", styles['Italic']))
    doc.build(elements)
    buffer.seek(0)
    return FileResponse(buffer, as_attachment=True, filename='filieres_dretfp.pdf')


@login_required
def export_filieres_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Filieres"
    headers = ['Code', 'Nom', 'Etablissement', 'Niveau', 'Nb apprenants', 'Duree', 'Actif']
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0a2540", end_color="0a2540", fill_type="solid")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    for row, f in enumerate(Filiere.objects.select_related('etablissement').all(), start=2):
        ws.cell(row=row, column=1, value=f.code)
        ws.cell(row=row, column=2, value=f.nom)
        ws.cell(row=row, column=3, value=f.etablissement.nom)
        ws.cell(row=row, column=4, value=f.get_niveau_display())
        ws.cell(row=row, column=5, value=f.nb_apprenants)
        ws.cell(row=row, column=6, value=f.duree_annees)
        ws.cell(row=row, column=7, value='Oui' if f.actif else 'Non')
    for i, w in enumerate([10, 30, 25, 12, 14, 10, 8], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="filieres_dretfp.xlsx"'
    wb.save(response)
    return response


@login_required
def export_salles_pdf(request):
    salles = Salle.objects.select_related('etablissement').all()
    etab_id = request.GET.get('etablissement')
    if etab_id:
        salles = salles.filter(etablissement_id=etab_id)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1*cm, leftMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    elements = []
    styles = getSampleStyleSheet()
    titre_style = ParagraphStyle('Titre', parent=styles['Heading1'], alignment=TA_CENTER,
        textColor=colors.HexColor('#0a2540'), fontSize=16, spaceAfter=10)
    elements.append(Paragraph("DRETFP Manager - Amoron'i Mania", titre_style))
    elements.append(Paragraph("Liste des salles", styles['Heading2']))
    elements.append(Spacer(1, 0.5*cm))
    data = [['Nom', 'Etablissement', 'Capacite', 'Type', 'Batiment', 'Actif']]
    for s in salles:
        data.append([s.nom[:30], s.etablissement.nom[:25], str(s.capacite),
                     s.type_salle[:20] if s.type_salle else '-',
                     s.batiment[:20] if s.batiment else '-', 'Oui' if s.actif else 'Non'])
    table = Table(data, colWidths=[4*cm, 4.5*cm, 1.8*cm, 2.5*cm, 2.5*cm, 1.5*cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a2540')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.5*cm))
    elements.append(Paragraph(f"Total : {salles.count()} salle(s)", styles['Normal']))
    elements.append(Paragraph(f"Genere le {datetime.now().strftime('%d/%m/%Y a %H:%M')}", styles['Italic']))
    doc.build(elements)
    buffer.seek(0)
    return FileResponse(buffer, as_attachment=True, filename='salles_dretfp.pdf')


@login_required
def export_matieres_pdf(request):
    matieres = Matiere.objects.select_related('filiere', 'filiere__etablissement').all()
    filiere_id = request.GET.get('filiere')
    if filiere_id:
        matieres = matieres.filter(filiere_id=filiere_id)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1*cm, leftMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    elements = []
    styles = getSampleStyleSheet()
    titre_style = ParagraphStyle('Titre', parent=styles['Heading1'], alignment=TA_CENTER,
        textColor=colors.HexColor('#0a2540'), fontSize=16, spaceAfter=10)
    elements.append(Paragraph("DRETFP Manager - Amoron'i Mania", titre_style))
    elements.append(Paragraph("Liste des matieres", styles['Heading2']))
    elements.append(Spacer(1, 0.5*cm))
    data = [['Code', 'Matiere', 'Filiere', 'Type', 'H.Theo', 'H.Prat', 'Total', 'Coef']]
    for m in matieres:
        data.append([m.code, m.nom[:35], m.filiere.code, m.type_matiere[:15],
                     str(m.heures_theorie), str(m.heures_pratique),
                     str(m.heures_total), str(m.coefficient)])
    table = Table(data, colWidths=[2*cm, 5*cm, 2*cm, 3*cm, 1.5*cm, 1.5*cm, 1.5*cm, 1.2*cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a2540')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.5*cm))
    elements.append(Paragraph(f"Total : {matieres.count()} matiere(s)", styles['Normal']))
    elements.append(Paragraph(f"Genere le {datetime.now().strftime('%d/%m/%Y a %H:%M')}", styles['Italic']))
    doc.build(elements)
    buffer.seek(0)
    return FileResponse(buffer, as_attachment=True, filename='matieres_dretfp.pdf')


@login_required
def export_formateurs_pdf(request):
    formateurs = Profil.objects.filter(role='FORMATEUR').select_related('utilisateur', 'etablissement')
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1*cm, leftMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    elements = []
    styles = getSampleStyleSheet()
    titre_style = ParagraphStyle('Titre', parent=styles['Heading1'], alignment=TA_CENTER,
        textColor=colors.HexColor('#0a2540'), fontSize=16, spaceAfter=10)
    elements.append(Paragraph("DRETFP Manager - Amoron'i Mania", titre_style))
    elements.append(Paragraph("Liste des formateurs", styles['Heading2']))
    elements.append(Spacer(1, 0.5*cm))
    data = [['Nom & Prenom', 'Username', 'Etablissement', 'Telephone', 'Email']]
    for f in formateurs:
        data.append([
            f"{f.utilisateur.last_name} {f.utilisateur.first_name}".strip() or '-',
            f.utilisateur.username,
            f.etablissement.nom[:25] if f.etablissement else '-',
            f.telephone or '-', f.utilisateur.email[:30] or '-'])
    table = Table(data, colWidths=[4*cm, 3*cm, 4.5*cm, 2.5*cm, 4*cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a2540')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.5*cm))
    elements.append(Paragraph(f"Total : {formateurs.count()} formateur(s)", styles['Normal']))
    elements.append(Paragraph(f"Genere le {datetime.now().strftime('%d/%m/%Y a %H:%M')}", styles['Italic']))
    doc.build(elements)
    buffer.seek(0)
    return FileResponse(buffer, as_attachment=True, filename='formateurs_dretfp.pdf')


@login_required
def export_emploi_pdf(request, filiere_id):
    filiere = get_object_or_404(Filiere, id=filiere_id)
    creneaux = CreneauHoraire.objects.all().order_by('ordre')
    jours = ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi']
    emplois = EmploiDuTemps.objects.filter(filiere=filiere).select_related('salle', 'creneau', 'formateur')
    grille = {}
    for e in emplois:
        grille[(e.jour, e.creneau.id)] = e
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1*cm, leftMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    elements = []
    styles = getSampleStyleSheet()
    titre_style = ParagraphStyle('Titre', parent=styles['Heading1'], alignment=TA_CENTER,
        textColor=colors.HexColor('#0a2540'), fontSize=16, spaceAfter=10)
    elements.append(Paragraph("DRETFP Manager - Amoron'i Mania", titre_style))
    elements.append(Paragraph("Emploi du temps", styles['Heading2']))
    elements.append(Paragraph(f"{filiere.code} - {filiere.nom}", styles['Normal']))
    elements.append(Paragraph(f"Etablissement : {filiere.etablissement.nom}", styles['Normal']))
    elements.append(Spacer(1, 0.5*cm))
    data = [['Horaire'] + jours]
    for c in creneaux:
        row = [f"{c.nom}\n{c.heure_debut.strftime('%H:%M')}-{c.heure_fin.strftime('%H:%M')}"]
        for j in jours:
            e = grille.get((j, c.id))
            if e:
                txt = f"{e.matiere}\n"
                if e.salle:
                    txt += f"Salle: {e.salle.nom}\n"
                if e.formateur:
                    txt += f"{e.formateur.get_full_name() or e.formateur.username}"
                row.append(txt)
            else:
                row.append('-')
        data.append(row)
    table = Table(data, colWidths=[2.5*cm, 2.8*cm, 2.8*cm, 2.8*cm, 2.8*cm, 2.8*cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a2540')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.5*cm))
    elements.append(Paragraph(f"Genere le {datetime.now().strftime('%d/%m/%Y a %H:%M')}", styles['Italic']))
    doc.build(elements)
    buffer.seek(0)
    nom_fichier = f"Emploi_du_temps_{filiere.code}.pdf".replace(' ', '_')
    return FileResponse(buffer, as_attachment=True, filename=nom_fichier)


# ============================================================
# EXPORTS BUDGETS
# ============================================================

@login_required
def export_budgets_csv(request):
    import csv
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="budgets_dretfp.csv"'
    response.write('\ufeff')
    writer = csv.writer(response, delimiter=';')
    writer.writerow(['Etablissement', 'Produit', 'Source', 'PCOP', 'Montant prevu (Ar)', 'Montant realise (Ar)', 'Statut'])
    for l in LigneBudgetaire.objects.select_related('etablissement', 'produit', 'pcop').all():
        writer.writerow([l.etablissement.nom, l.produit.libelle, l.source_financement,
                         l.pcop.code if l.pcop else '', l.montant_prevu, l.montant_realise, l.statut])
    return response


@login_required
def export_budgets_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Budgets DRETFP"
    headers = ['Etablissement', 'Produit', 'Source', 'PCOP', 'Montant prevu (Ar)', 'Montant realise (Ar)', 'Statut']
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="0a2540", end_color="0a2540", fill_type="solid")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    for row, l in enumerate(LigneBudgetaire.objects.select_related('etablissement', 'produit', 'pcop').all(), start=2):
        ws.cell(row=row, column=1, value=l.etablissement.nom)
        ws.cell(row=row, column=2, value=l.produit.libelle)
        ws.cell(row=row, column=3, value=l.source_financement)
        ws.cell(row=row, column=4, value=l.pcop.code if l.pcop else '')
        ws.cell(row=row, column=5, value=float(l.montant_prevu))
        ws.cell(row=row, column=6, value=float(l.montant_realise))
        ws.cell(row=row, column=7, value=l.statut)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="budgets_dretfp.xlsx"'
    wb.save(response)
    return response


# ============================================================
# RAPPORTS
# ============================================================

@login_required
def rapports(request):
    return render(request, 'core/rapports.html', {'etablissements': Etablissement.objects.all()})


@login_required
def rapport_pdf_global(request):
    etab_id = request.GET.get('etablissement')
    source = request.GET.get('source')
    lignes = LigneBudgetaire.objects.select_related('etablissement', 'produit', 'pcop').all()
    if etab_id:
        lignes = lignes.filter(etablissement_id=etab_id)
    if source:
        lignes = lignes.filter(source_financement=source)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1*cm, leftMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    elements = []
    styles = getSampleStyleSheet()
    titre_style = ParagraphStyle('Titre', parent=styles['Heading1'], alignment=TA_CENTER,
        textColor=colors.HexColor('#0a2540'), fontSize=18, spaceAfter=12)
    elements.append(Paragraph("DRETFP Manager - Amoron'i Mania", titre_style))
    elements.append(Paragraph("Rapport des lignes budgetaires", styles['Heading2']))
    elements.append(Spacer(1, 0.5*cm))
    total_prevu = lignes.aggregate(Sum('montant_prevu'))['montant_prevu__sum'] or 0
    total_realise = lignes.aggregate(Sum('montant_realise'))['montant_realise__sum'] or 0
    kpi_data = [
        ['Nombre de lignes', str(lignes.count())],
        ['Montant total prevu', f"{total_prevu:,.0f} Ar"],
        ['Montant total realise', f"{total_realise:,.0f} Ar"],
    ]
    kpi_table = Table(kpi_data, colWidths=[8*cm, 8*cm])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#e7f1ff')),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 0.5*cm))
    data = [['#', 'Etablissement', 'Produit', 'Src', 'Montant (Ar)', 'Statut']]
    for i, l in enumerate(lignes[:500], start=1):
        data.append([str(i), l.etablissement.nom[:20], l.produit.libelle[:40],
                     l.source_financement, f"{float(l.montant_prevu):,.0f}", l.statut])
    table = Table(data, colWidths=[0.8*cm, 3*cm, 6*cm, 1.2*cm, 3*cm, 2*cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a2540')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.5*cm))
    elements.append(Paragraph(f"Genere le {datetime.now().strftime('%d/%m/%Y a %H:%M')}", styles['Italic']))
    doc.build(elements)
    buffer.seek(0)
    return FileResponse(buffer, as_attachment=True, filename='rapport_budgets_dretfp.pdf')


# ============================================================
# INSCRIPTION
# ============================================================

class InscriptionForm(UserCreationForm):
    email = forms.EmailField(required=True, label="Email")
    first_name = forms.CharField(max_length=30, required=True, label="Prenom")
    last_name = forms.CharField(max_length=30, required=True, label="Nom")
    
    class Meta:
        model = User
        fields = ('username', 'first_name', 'last_name', 'email', 'password1', 'password2')
    
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        for field in self.fields.values():
            field.widget.attrs.update({'class': 'form-control'})


def inscription(request):
    if request.method == 'POST':
        form = InscriptionForm(request.POST)
        if form.is_valid():
            user = form.save()
            Profil.objects.create(utilisateur=user, role='LECTEUR')
            login(request, user)
            messages.success(request, "Bienvenue " + user.username + " !")
            return redirect('accueil')
    else:
        form = InscriptionForm()
    return render(request, 'registration/inscription.html', {'form': form})


# ============================================================
# MOT DE PASSE OUBLIE
# ============================================================

def mot_de_passe_oublie(request):
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        email_contact = request.POST.get('email_contact', '').strip()
        message = request.POST.get('message', '').strip()
        if not username:
            messages.error(request, "Veuillez entrer votre nom d'utilisateur.")
        else:
            try:
                user = User.objects.get(username=username)
                existante = DemandeReinitialisation.objects.filter(utilisateur=user, statut='En attente').exists()
                if existante:
                    messages.warning(request, "Une demande est deja en attente.")
                else:
                    DemandeReinitialisation.objects.create(utilisateur=user, email_contact=email_contact, message=message)
                    messages.success(request, "Votre demande a ete envoyee.")
                    return redirect('login')
            except User.DoesNotExist:
                messages.error(request, "Aucun compte trouve.")
    return render(request, 'registration/mot_de_passe_oublie.html')


@login_required
def liste_demandes(request):
    if not request.user.is_staff:
        messages.error(request, "Acces reserve aux administrateurs.")
        return redirect('accueil')
    demandes = DemandeReinitialisation.objects.all()
    context = {'demandes': demandes, 'nb_en_attente': demandes.filter(statut='En attente').count()}
    return render(request, 'core/liste_demandes.html', context)


@login_required
def traiter_demande(request, demande_id):
    if not request.user.is_staff:
        messages.error(request, "Acces reserve aux administrateurs.")
        return redirect('accueil')
    demande = get_object_or_404(DemandeReinitialisation, id=demande_id)
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'reinitialiser':
            from django.utils.crypto import get_random_string
            from django.utils import timezone
            nouveau_mdp = 'Dretfp' + get_random_string(6)
            user = demande.utilisateur
            user.set_password(nouveau_mdp)
            user.save()
            demande.nouveau_mot_de_passe = nouveau_mdp
            demande.statut = 'Traitee'
            demande.date_traitement = timezone.now()
            demande.save()
            messages.success(request, "Mot de passe reinitialise pour " + user.username + " ! Nouveau : " + nouveau_mdp)
            return redirect('liste_demandes')
        elif action == 'rejeter':
            from django.utils import timezone
            demande.statut = 'Rejetee'
            demande.date_traitement = timezone.now()
            demande.save()
            messages.info(request, "Demande rejetee.")
            return redirect('liste_demandes')
    return render(request, 'core/traiter_demande.html', {'demande': demande})


# ============================================================
# API JSON
# ============================================================

def api_sources(request):
    data = LigneBudgetaire.objects.values('source_financement').annotate(total=Sum('montant_prevu')).order_by('-total')
    return JsonResponse({'labels': [d['source_financement'] for d in data], 'values': [float(d['total'] or 0) for d in data]})

def api_etablissements(request):
    data = LigneBudgetaire.objects.values('etablissement__nom').annotate(total=Sum('montant_prevu')).order_by('-total')
    return JsonResponse({'labels': [d['etablissement__nom'] for d in data], 'values': [float(d['total'] or 0) for d in data]})

def api_produits(request):
    data = LigneBudgetaire.objects.values('produit__libelle').annotate(total=Sum('montant_prevu')).order_by('-total')[:10]
    return JsonResponse({'labels': [d['produit__libelle'] for d in data], 'values': [float(d['total'] or 0) for d in data]})

def api_statuts(request):
    data = LigneBudgetaire.objects.values('statut').annotate(count=Count('id')).order_by('statut')
    return JsonResponse({'labels': [d['statut'] for d in data], 'values': [d['count'] for d in data]})

def api_top_etablissements(request):
    data = LigneBudgetaire.objects.values('etablissement__nom').annotate(total=Sum('montant_prevu')).order_by('-total')[:10]
    return JsonResponse({'labels': [d['etablissement__nom'] for d in data], 'values': [float(d['total'] or 0) for d in data]})

def api_types_etablissements(request):
    data = Etablissement.objects.values('type_etablissement').annotate(count=Count('id')).order_by('type_etablissement')
    types_labels = {'LTP': 'Lycee Technique', 'CFP': 'Centre de Formation', 'CFPF': 'CFP + Formation', 'LTPA': 'Lycee Agricole'}
    return JsonResponse({'labels': [types_labels.get(d['type_etablissement'], d['type_etablissement']) for d in data], 'values': [d['count'] for d in data]})

def api_rpi_fce_par_etab(request):
    etablissements = Etablissement.objects.all()
    labels, rpi_values, fce_values = [], [], []
    for etab in etablissements:
        labels.append(etab.nom)
        rpi = LigneBudgetaire.objects.filter(etablissement=etab, source_financement='RPI').aggregate(Sum('montant_prevu'))['montant_prevu__sum'] or 0
        fce = LigneBudgetaire.objects.filter(etablissement=etab, source_financement='FCE').aggregate(Sum('montant_prevu'))['montant_prevu__sum'] or 0
        rpi_values.append(float(rpi))
        fce_values.append(float(fce))
    return JsonResponse({'labels': labels, 'rpi': rpi_values, 'fce': fce_values})

def api_top_etablissements_montant(request):
    data = LigneBudgetaire.objects.values('etablissement__nom').annotate(total=Sum('montant_prevu')).order_by('-total')[:10]
    return JsonResponse({'labels': [d['etablissement__nom'] for d in data], 'values': [float(d['total'] or 0) for d in data]})

def api_produits_par_etab(request):
    etablissements = Etablissement.objects.all()
    labels, values = [], []
    for etab in etablissements:
        count = LigneBudgetaire.objects.filter(etablissement=etab).values('produit').distinct().count()
        labels.append(etab.nom)
        values.append(count)
    return JsonResponse({'labels': labels, 'values': values})

def api_taux_execution_etab(request):
    etablissements = Etablissement.objects.all()
    labels, values = [], []
    for etab in etablissements:
        total_prevu = LigneBudgetaire.objects.filter(etablissement=etab).aggregate(Sum('montant_prevu'))['montant_prevu__sum'] or 0
        total_realise = LigneBudgetaire.objects.filter(etablissement=etab).aggregate(Sum('montant_realise'))['montant_realise__sum'] or 0
        taux = 0
        if total_prevu > 0:
            taux = round((float(total_realise) / float(total_prevu)) * 100, 2)
        labels.append(etab.nom)
        values.append(taux)
    return JsonResponse({'labels': labels, 'values': values})

def api_statuts_detail(request):
    data = LigneBudgetaire.objects.values('statut').annotate(count=Count('id'), total=Sum('montant_prevu')).order_by('statut')
    return JsonResponse({
        'labels': [d['statut'] for d in data],
        'counts': [d['count'] for d in data],
        'totals': [float(d['total'] or 0) for d in data],
    })